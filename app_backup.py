
import io
import re
import streamlit as st
from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ==================================================
# 1. PAGE SETTINGS
# ==================================================

st.set_page_config(
    page_title="CareerForge AI",
    page_icon="🚀",
    layout="wide"
)

st.title("🚀 CareerForge AI")
st.subheader("CV Matcher & Career Intelligence")

st.write(
    "Upload your CV and paste a job description to identify skills, "
    "compare your CV with the job, review possible skill gaps, "
    "and generate editable resume content."
)


# ==================================================
# 2. SKILL TAXONOMY
# ==================================================

SKILL_TAXONOMY = {
    "Programming Languages": [
        "Python", "Java", "C++", "JavaScript",
        "TypeScript", "Go", "Rust", "SQL"
    ],
    "Data & AI / ML": [
        "Pandas", "NumPy", "scikit-learn", "TensorFlow",
        "PyTorch", "FastAPI", "OpenCV", "NLP",
        "Data Science", "Machine Learning"
    ],
    "Frameworks & Web": [
        "Streamlit", "Django", "Flask",
        "React", "Angular", "Node.js"
    ],
    "Tools & Platforms": [
        "Git", "GitHub", "Docker", "Kubernetes",
        "AWS", "Azure", "Linux", "Postman",
        "Zapier", "Cisco Packet Tracer"
    ],
}


# ==================================================
# 3. SKILL ALIASES
# ==================================================

SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "C++": ["c++", "cpp"],
    "JavaScript": ["javascript", "java script"],
    "TypeScript": ["typescript", "type script"],
    "Go": ["golang", "go"],
    "Rust": ["rust"],
    "SQL": ["sql", "structured query language"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy", "num py"],
    "scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "TensorFlow": ["tensorflow", "tensor flow"],
    "PyTorch": ["pytorch", "py torch"],
    "FastAPI": ["fastapi", "fast api"],
    "OpenCV": ["opencv", "open cv"],
    "NLP": ["nlp", "natural language processing"],
    "Data Science": ["data science"],
    "Machine Learning": ["machine learning"],
    "Streamlit": ["streamlit"],
    "Django": ["django"],
    "Flask": ["flask"],
    "React": ["react"],
    "Angular": ["angular"],
    "Node.js": ["node.js", "node js"],
    "Git": ["git"],
    "GitHub": ["github", "git hub"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure", "microsoft azure"],
    "Linux": ["linux"],
    "Postman": ["postman"],
    "Zapier": ["zapier"],
    "Cisco Packet Tracer": ["cisco packet tracer", "packet tracer"],
}


# ==================================================
# 4. FILE TEXT EXTRACTION
# ==================================================

def extract_pdf_text(uploaded_file):
    try:
        reader = PdfReader(io.BytesIO(uploaded_file.getvalue()))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )
    except Exception as error:
        st.error(f"Could not read PDF: {error}")
        return ""


def extract_docx_text(uploaded_file):
    try:
        document = Document(io.BytesIO(uploaded_file.getvalue()))

        text_parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        return "\n".join(text_parts)

    except Exception as error:
        st.error(f"Could not read Word document: {error}")
        return ""


def extract_txt_text(uploaded_file):
    try:
        return uploaded_file.getvalue().decode(
            "utf-8", errors="replace"
        )
    except Exception as error:
        st.error(f"Could not read text file: {error}")
        return ""


def extract_file_text(uploaded_file):
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    elif file_name.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    elif file_name.endswith(".txt"):
        return extract_txt_text(uploaded_file)

    return ""


# ==================================================
# 5. SKILL EXTRACTION
# ==================================================

def skill_is_present(text, alias):
    """
    Match a skill as a separate word or phrase.
    Avoid false matches such as 'go' inside 'ongoing'.
    """
    text = text.lower()
    alias = alias.lower()

    pattern = (
        r"(?<![a-z0-9])"
        + re.escape(alias)
        + r"(?![a-z0-9])"
    )

    return re.search(pattern, text) is not None


def extract_all_skills(text):
    found_skills = set()
    skills_by_category = {}

    if not text or not text.strip():
        return skills_by_category, found_skills

    for category, skills in SKILL_TAXONOMY.items():
        category_matches = []

        for skill in skills:
            aliases = SKILL_ALIASES.get(
                skill, [skill.lower()]
            )

            if any(
                skill_is_present(text, alias)
                for alias in aliases
            ):
                category_matches.append(skill)
                found_skills.add(skill)

        if category_matches:
            skills_by_category[category] = category_matches

    return skills_by_category, found_skills


# ==================================================
# 6. SEPARATE REQUIRED AND PREFERRED SKILLS
# ==================================================

def extract_requirement_sections(job_description):
    """
    Separate skills appearing before the Preferred Skills heading
    from skills appearing in the preferred section.

    This works best when the job description has clear headings.
    """

    lower_text = job_description.lower()

    preferred_heading = re.search(
        r"(?im)^\s*(?:#{1,6}\s*)?"
        r"(?:preferred skills|preferred qualifications|"
        r"nice[- ]to[- ]have skills)\s*:?\s*$",
        lower_text
    )

    if preferred_heading:
        required_text = job_description[
            :preferred_heading.start()
        ]

        preferred_text = job_description[
            preferred_heading.end():
        ]

        # Stop preferred text at another major section heading.
        next_heading = re.search(
            r"(?im)^\s*(?:#{1,6}\s*)?"
            r"(?:what we offer|benefits|about us|"
            r"how to apply|application process)\s*:?\s*$",
            preferred_text
        )

        if next_heading:
            preferred_text = preferred_text[
                :next_heading.start()
            ]

    else:
        # If no preferred section exists, treat all detected
        # job-description skills as requirements to review.
        required_text = job_description
        preferred_text = ""

    _, all_required_skills = extract_all_skills(required_text)
    _, preferred_skills = extract_all_skills(preferred_text)

    # Do not count preferred skills as required.
    required_skills = all_required_skills - preferred_skills

    return required_skills, preferred_skills


# ==================================================
# 7. TF-IDF TEXT SIMILARITY
# ==================================================

def calculate_text_similarity(cv_text, job_text):
    if not cv_text.strip() or not job_text.strip():
        return 0.0

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2)
        )

        vectors = vectorizer.fit_transform(
            [cv_text, job_text]
        )

        score = cosine_similarity(
            vectors[0:1],
            vectors[1:2]
        )[0][0]

        return round(float(score) * 100, 2)

    except ValueError:
        return 0.0


# ==================================================
# 8. RESUME CONTENT GENERATION
# ==================================================

def generate_resume_content(
    cv_skills,
    matched_required,
    missing_required,
    matched_preferred,
    missing_preferred,
    similarity_score
):
    if cv_skills:
        cv_skill_text = ", ".join(sorted(cv_skills))
    else:
        cv_skill_text = "No taxonomy skills confidently detected"

    if matched_required:
        matched_text = ", ".join(sorted(matched_required))
    else:
        matched_text = "No matching required skills detected"

    if missing_required:
        missing_required_text = ", ".join(
            sorted(missing_required)
        )
    else:
        missing_required_text = "None detected"

    if missing_preferred:
        missing_preferred_text = ", ".join(
            sorted(missing_preferred)
        )
    else:
        missing_preferred_text = "None detected"

    summary = (
        "Software Engineering candidate with academic project "
        "experience reflected in the uploaded CV. Skills detected "
        f"by the current extractor include: {cv_skill_text}. "
        "Verify every skill and edit this summary before using it."
    )

    bullets = [
        (
            "Skills detected in the uploaded CV: "
            f"{cv_skill_text}. Keep only skills that accurately "
            "represent your abilities."
        ),
        (
            "Required skills detected in both the CV and job "
            f"description: {matched_text}."
        ),
        (
            "Required skills not detected in the CV: "
            f"{missing_required_text}. Review these items; "
            "they are not proof that you lack the skills."
        ),
        (
            "Preferred skills not detected in the CV: "
            f"{missing_preferred_text}. Consider these optional "
            "learning areas if they are relevant to your goals."
        ),
        (
            "TF-IDF text similarity: "
            f"{similarity_score}%. This measures text overlap "
            "and is not a direct measure of job suitability."
        ),
    ]

    return summary, bullets


# ==================================================
# 9. CREATE WORD DOCUMENT
# ==================================================

def create_docx_file(summary_text, bullet_points):
    document = Document()

    document.add_heading(
        "CareerForge AI - Resume Content", level=1
    )

    document.add_heading(
        "Professional Summary Draft", level=2
    )
    document.add_paragraph(summary_text)

    document.add_heading(
        "Skills and Review Notes", level=2
    )

    for bullet in bullet_points:
        document.add_paragraph(
            bullet, style="List Bullet"
        )

    document.add_paragraph(
        "Review all generated content before using it in a CV."
    )

    file_buffer = io.BytesIO()
    document.save(file_buffer)
    file_buffer.seek(0)

    return file_buffer.getvalue()


# ==================================================
# 10. SIDEBAR
# ==================================================

with st.sidebar:
    st.header("About CareerForge AI")

    st.write(
        "CareerForge AI compares CVs with job descriptions "
        "using skill extraction and TF-IDF text similarity."
    )

    st.caption(
        "This is an academic prototype. Automatically detected "
        "skills and scores should be reviewed."
    )


# ==================================================
# 11. UPLOAD CV
# ==================================================

st.header("1. Upload Candidate Resume")

uploaded_file = st.file_uploader(
    "Choose your CV or resume",
    type=["pdf", "docx", "txt"]
)

cv_text = ""
cv_skills_by_category = {}
cv_skills = set()

if uploaded_file is not None:
    cv_text = extract_file_text(uploaded_file)

    if cv_text.strip():
        st.success("CV text extracted successfully.")

        with st.expander("Preview extracted CV text"):
            st.text(cv_text[:10000])

        cv_skills_by_category, cv_skills = extract_all_skills(
            cv_text
        )

        st.subheader("Skills Detected in Your CV")

        if cv_skills_by_category:
            for category, skills in cv_skills_by_category.items():
                st.markdown(f"**{category}**")
                st.write(", ".join(skills))
        else:
            st.info("No skills from the current taxonomy were detected.")

        st.caption(
            "Please verify detected skills against your original CV."
        )

    else:
        st.warning(
            "No readable text was extracted from the uploaded file."
        )


# ==================================================
# 12. JOB DESCRIPTION
# ==================================================

st.header("2. Target Job Description")

job_description = st.text_area(
    "Paste the job description here",
    height=250,
    placeholder="Paste the complete job description here..."
)


# ==================================================
# 13. ANALYSIS
# ==================================================

if st.button("Analyze CV and Job Description", type="primary"):

    if uploaded_file is None:
        st.error("Please upload your CV first.")

    elif not cv_text.strip():
        st.error("No readable CV text was found.")

    elif not job_description.strip():
        st.error("Please paste a job description.")

    else:
        # Separate required and preferred skills.
        required_skills, preferred_skills = (
            extract_requirement_sections(job_description)
        )

        # Compare CV skills with required skills.
        matched_required = cv_skills.intersection(
            required_skills
        )
        missing_required = required_skills.difference(
            cv_skills
        )

        # Compare CV skills with preferred skills.
        matched_preferred = cv_skills.intersection(
            preferred_skills
        )
        missing_preferred = preferred_skills.difference(
            cv_skills
        )

        # Combine all matches for overall display.
        matched_skills = (
            matched_required.union(matched_preferred)
        )
        missing_skills = (
            missing_required.union(missing_preferred)
        )

        # Score required-skill coverage only.
        if required_skills:
            taxonomy_score = round(
                len(matched_required)
                / len(required_skills)
                * 100
            )
        else:
            taxonomy_score = 0

        # Calculate TF-IDF text similarity.
        similarity_score = calculate_text_similarity(
            cv_text,
            job_description
        )

        # Generate resume content.
        summary_text, bullet_points = generate_resume_content(
            cv_skills,
            matched_required,
            missing_required,
            matched_preferred,
            missing_preferred,
            similarity_score
        )

        # Save results across Streamlit reruns.
        st.session_state["analysis_result"] = {
            "cv_skills": sorted(cv_skills),
            "required_skills": sorted(required_skills),
            "preferred_skills": sorted(preferred_skills),
            "matched_required": sorted(matched_required),
            "missing_required": sorted(missing_required),
            "matched_preferred": sorted(matched_preferred),
            "missing_preferred": sorted(missing_preferred),
            "matched_skills": sorted(matched_skills),
            "missing_skills": sorted(missing_skills),
            "taxonomy_score": taxonomy_score,
            "similarity_score": similarity_score,
            "summary_text": summary_text,
            "bullet_points": bullet_points,
        }


# ==================================================
# 14. DISPLAY RESULTS
# ==================================================

if "analysis_result" in st.session_state:
    result = st.session_state["analysis_result"]

    st.divider()
    st.header("3. Analysis Results")

    col1, col2 = st.columns(2)

    with col1:
        st.metric(
            "TF-IDF Text Similarity",
            f"{result['similarity_score']}%"
        )

    with col2:
        st.metric(
            "Required-Skill Coverage",
            f"{result['taxonomy_score']}%"
        )

    st.progress(result["taxonomy_score"] / 100)

    st.caption(
        "Required-skill coverage is based only on skills detected "
        "by the current taxonomy. It does not measure every job "
        "requirement or predict your chance of being hired."
    )

    # Required skills section.
    st.subheader("Required Skills")

    st.markdown("**Matched required skills**")
    if result["matched_required"]:
        st.write(", ".join(result["matched_required"]))
    else:
        st.info("No required skills matched.")

    st.markdown("**Required skills to review**")
    if result["missing_required"]:
        st.write(", ".join(result["missing_required"]))
    else:
        st.success("All detected required skills matched.")

    # Preferred skills section.
    st.subheader("Preferred Skills")

    st.markdown("**Matched preferred skills**")
    if result["matched_preferred"]:
        st.write(", ".join(result["matched_preferred"]))
    else:
        st.write("No preferred skills matched.")

    st.markdown("**Preferred skills to consider learning**")
    if result["missing_preferred"]:
        st.write(", ".join(result["missing_preferred"]))
    else:
        st.write("All detected preferred skills matched.")

    # View detected skills.
    with st.expander("View all detected job-description skills"):
        st.markdown("**Required skills**")
        st.write(
            ", ".join(result["required_skills"])
            if result["required_skills"]
            else "None detected"
        )

        st.markdown("**Preferred skills**")
        st.write(
            ", ".join(result["preferred_skills"])
            if result["preferred_skills"]
            else "None detected"
        )

    # Resume content.
    st.divider()
    st.header("4. Generated Resume Content")

    st.subheader("Professional Summary Draft")
    st.write(result["summary_text"])

    st.subheader("Resume Bullet Points")
    for bullet in result["bullet_points"]:
        st.markdown(f"- {bullet}")

    st.warning(
        "Check the generated content before using it in your CV. "
        "Only include skills and experience you genuinely possess."
    )

    # TXT content.
    text_content = (
        "CAREERFORGE AI - RESUME CONTENT\n\n"
        "PROFESSIONAL SUMMARY\n"
        f"{result['summary_text']}\n\n"
        "SKILLS AND REVIEW NOTES\n"
        + "\n".join(
            f"- {bullet}"
            for bullet in result["bullet_points"]
        )
    )

    # Word document content.
    docx_content = create_docx_file(
        result["summary_text"],
        result["bullet_points"]
    )

    download_col1, download_col2 = st.columns(2)

    with download_col1:
        st.download_button(
            label="Download Word Document",
            data=docx_content,
            file_name="CareerForge_Resume_Content.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

    with download_col2:
        st.download_button(
            label="Download Text File",
            data=text_content,
            file_name="CareerForge_Resume_Content.txt",
            mime="text/plain",
        )


# ==================================================
# 15. FOOTER
# ==================================================

st.divider()
st.caption("CareerForge AI | Academic Project Prototype")
