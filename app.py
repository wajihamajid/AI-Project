
import io
import re

import streamlit as st
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from app.document_processor import extract_text


# ==================================================
# 1. PAGE SETTINGS
# ==================================================

st.set_page_config(
    page_title="CareerForge AI",
    page_icon="🚀",
    layout="wide"
)

st.title("🚀 CareerForge AI")
st.subheader("CV Matcher & Career Intelligence")

st.write(
    "Upload a CV and paste a job description to identify skills, "
    "compare relevant evidence, review potential skill gaps, "
    "and generate editable resume content."
)


# ==================================================
# 2. GENERAL-PURPOSE SKILL TAXONOMY
# ==================================================

SKILL_TAXONOMY = {
    "Programming Languages": [
        "Python", "Java", "C++", "C#", "JavaScript",
        "TypeScript", "Go", "Rust", "PHP", "Ruby",
        "Swift", "Kotlin", "R", "SQL", "Bash",
        "PowerShell"
    ],

    "Data Preparation & Analytics": [
        "Pandas", "NumPy", "Data Cleaning",
        "Data Preprocessing", "Data Encoding",
        "Feature Scaling", "Feature Engineering",
        "Data Visualization", "Statistics",
        "Excel", "Power BI", "Tableau", "ETL"
    ],

    "AI & Machine Learning": [
        "scikit-learn", "TensorFlow", "PyTorch",
        "OpenCV", "NLP", "Data Science",
        "Machine Learning", "Deep Learning",
        "Generative AI", "Large Language Models",
        "Computer Vision", "Natural Language Processing"
    ],

    "Web Development & Frameworks": [
        "Streamlit", "Django", "Flask", "FastAPI",
        "React", "Angular", "Node.js", "Express.js",
        "HTML", "CSS", "REST API"
    ],

    "Databases": [
        "PostgreSQL", "MySQL", "SQLite",
        "MongoDB", "Redis", "Database Design"
    ],

    "Software Engineering": [
        "Object-Oriented Programming",
        "Software Design", "UML",
        "Requirements Analysis", "Software Testing",
        "Debugging", "Software Documentation",
        "Version Control", "Agile", "Scrum",
        "Software Development Life Cycle"
    ],

    "Cloud & DevOps": [
        "Git", "GitHub", "Docker", "Kubernetes",
        "AWS", "Azure", "Google Cloud",
        "Linux", "CI/CD", "Jenkins", "Terraform"
    ],

    "Networking & Security": [
        "Cisco Packet Tracer", "Networking",
        "TCP/IP", "VLAN", "Routing", "OSPF",
        "NAT", "Port Security", "DNS", "DHCP",
        "Network Security", "Firewall"
    ],

    "Tools & Productivity": [
        "Postman", "Zapier", "VS Code",
        "Jira", "Microsoft Office"
    ],
}


# ==================================================
# 3. SKILL ALIASES
# ==================================================

# Aliases help recognize common names and variations.
# A detected phrase is not proof of proficiency.

SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "c sharp"],
    "JavaScript": ["javascript", "java script"],
    "TypeScript": ["typescript", "type script"],
    "Go": ["golang", "go programming", "go language"],
    "Rust": ["rust"],
    "PHP": ["php"],
    "Ruby": ["ruby"],
    "Swift": ["swift programming", "swift language"],
    "Kotlin": ["kotlin"],
    "R": ["r programming", "r language"],
    "SQL": ["sql", "structured query language"],
    "Bash": ["bash", "bash scripting"],
    "PowerShell": ["powershell", "power shell"],

    "Pandas": ["pandas"],
    "NumPy": ["numpy", "num py"],
    "Data Cleaning": [
        "data cleaning", "cleaning data", "cleaned datasets"
    ],
    "Data Preprocessing": [
        "data preprocessing", "data pre-processing",
        "preprocessing", "pre-processing"
    ],
    "Data Encoding": [
        "data encoding", "categorical encoding",
        "encoding categories", "encoding categorical data"
    ],
    "Feature Scaling": [
        "feature scaling", "scaling features",
        "feature normalization", "standardization of features"
    ],
    "Feature Engineering": ["feature engineering"],
    "Data Visualization": [
        "data visualization", "data visualisation",
        "data charts", "data visualization techniques"
    ],
    "Statistics": ["statistics", "statistical analysis"],
    "Excel": ["excel", "microsoft excel"],
    "Power BI": ["power bi", "powerbi"],
    "Tableau": ["tableau"],
    "ETL": ["etl", "extract transform load"],

    "scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "TensorFlow": ["tensorflow", "tensor flow"],
    "PyTorch": ["pytorch", "py torch"],
    "OpenCV": ["opencv", "open cv"],
    "NLP": ["nlp", "natural language processing"],
    "Data Science": ["data science"],
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "Generative AI": ["generative ai", "gen ai"],
    "Large Language Models": [
        "large language models", "large language model", "llms", "llm"
    ],
    "Computer Vision": ["computer vision"],
    "Natural Language Processing": [
        "natural language processing", "nlp"
    ],

    "Streamlit": ["streamlit"],
    "Django": ["django"],
    "Flask": ["flask"],
    "FastAPI": ["fastapi", "fast api"],
    "React": ["react", "react.js", "reactjs"],
    "Angular": ["angular"],
    "Node.js": ["node.js", "node js", "nodejs"],
    "Express.js": ["express.js", "express js", "expressjs"],
    "HTML": ["html", "html5"],
    "CSS": ["css", "css3"],
    "REST API": [
        "rest api", "rest apis", "restful api", "restful apis"
    ],

    "PostgreSQL": ["postgresql", "postgres"],
    "MySQL": ["mysql"],
    "SQLite": ["sqlite"],
    "MongoDB": ["mongodb", "mongo db"],
    "Redis": ["redis"],
    "Database Design": ["database design", "database modelling"],

    "Object-Oriented Programming": [
        "object-oriented programming", "object oriented programming",
        "oop", "object-oriented concepts"
    ],
    "Software Design": ["software design", "system design"],
    "UML": ["uml", "unified modeling language", "unified modelling language"],
    "Requirements Analysis": [
        "requirements analysis", "requirement analysis"
    ],
    "Software Testing": [
        "software testing", "unit testing", "integration testing"
    ],
    "Debugging": ["debugging", "debug code"],
    "Software Documentation": [
        "software documentation", "technical documentation",
        "project documentation", "code documentation"
    ],
    "Version Control": ["version control"],
    "Agile": ["agile methodology", "agile"],
    "Scrum": ["scrum"],
    "Software Development Life Cycle": [
        "software development life cycle", "sdlc"
    ],

    "Git": ["git"],
    "GitHub": ["github", "git hub"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes", "k8s"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure", "microsoft azure"],
    "Google Cloud": ["google cloud", "gcp"],
    "Linux": ["linux"],
    "CI/CD": ["ci/cd", "continuous integration", "continuous delivery"],
    "Jenkins": ["jenkins"],
    "Terraform": ["terraform"],

    "Cisco Packet Tracer": ["cisco packet tracer", "packet tracer"],
    "Networking": ["computer networking", "networking"],
    "TCP/IP": ["tcp/ip", "tcp ip"],
    "VLAN": ["vlan", "vlans", "virtual lan"],
    "Routing": ["routing", "network routing"],
    "OSPF": ["ospf", "open shortest path first"],
    "NAT": ["nat", "network address translation"],
    "Port Security": ["port security", "switch port security"],
    "DNS": ["dns", "domain name system"],
    "DHCP": ["dhcp", "dynamic host configuration protocol"],
    "Network Security": ["network security"],
    "Firewall": ["firewall", "firewalls"],

    "Postman": ["postman"],
    "Zapier": ["zapier"],
    "VS Code": ["vs code", "visual studio code"],
    "Jira": ["jira"],
    "Microsoft Office": ["microsoft office", "ms office"],
}


# ==================================================
# 4. SKILL EXTRACTION
# ==================================================

def skill_is_present(text, alias):
    """Find a phrase without matching it inside a longer word."""

    if not text or not alias:
        return False

    pattern = (
        r"(?<![a-z0-9])"
        + re.escape(alias.lower())
        + r"(?![a-z0-9])"
    )

    return re.search(pattern, text.lower()) is not None


def extract_all_skills(text):
    """Return detected taxonomy skills grouped by category."""

    found_skills = set()
    skills_by_category = {}

    if not text or not text.strip():
        return skills_by_category, found_skills

    for category, skills in SKILL_TAXONOMY.items():
        category_matches = []

        for skill in skills:
            aliases = SKILL_ALIASES.get(skill, [skill.lower()])

            if any(skill_is_present(text, alias) for alias in aliases):
                category_matches.append(skill)
                found_skills.add(skill)

        if category_matches:
            skills_by_category[category] = category_matches

    return skills_by_category, found_skills


# ==================================================
# 5. REQUIRED AND PREFERRED SKILLS
# ==================================================

def extract_requirement_sections(job_description):
    """
    Extract skills from the required portion and preferred portion.

    If a skill appears in both portions, required takes priority.
    When a JD has no preferred heading, detected skills are treated
    as requirements for this prototype.
    """

    if not job_description or not job_description.strip():
        return set(), set()

    preferred_heading = re.search(
        r"(?im)^\s*(?:#{1,6}\s*)?"
        r"(?:preferred skills|preferred qualifications|"
        r"preferred requirements|nice[- ]to[- ]have skills|"
        r"nice[- ]to[- ]have|bonus skills)\s*:?\s*$",
        job_description
    )

    if preferred_heading:
        required_text = job_description[:preferred_heading.start()]
        preferred_text = job_description[preferred_heading.end():]

        stop_heading = re.search(
            r"(?im)^\s*(?:#{1,6}\s*)?"
            r"(?:responsibilities|what you will do|"
            r"what we offer|benefits|about us|"
            r"how to apply|application process|"
            r"education|equal opportunity)\s*:?\s*$",
            preferred_text
        )

        if stop_heading:
            preferred_text = preferred_text[:stop_heading.start()]

    else:
        required_text = job_description
        preferred_text = ""

    _, required_skills = extract_all_skills(required_text)
    _, preferred_skills = extract_all_skills(preferred_text)

    # A skill explicitly found in required text remains required.
    # Avoid counting it a second time as preferred.
    preferred_skills = preferred_skills - required_skills

    return required_skills, preferred_skills


# ==================================================
# 6. TF-IDF TEXT SIMILARITY
# ==================================================

def calculate_text_similarity(cv_text, job_text):
    """Calculate text overlap, not a hiring probability."""

    if not cv_text.strip() or not job_text.strip():
        return 0.0

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2)
        )

        vectors = vectorizer.fit_transform([cv_text, job_text])

        score = cosine_similarity(
            vectors[0:1],
            vectors[1:2]
        )[0][0]

        return round(float(score) * 100, 2)

    except ValueError:
        return 0.0


# ==================================================
# 7. RESUME CONTENT GENERATION
# ==================================================


def generate_resume_content(
    cv_text,
    cv_skills,
    matched_required,
    missing_required,
    matched_preferred,
    missing_preferred,
    similarity_score
):
    cv_text_clean = " ".join(cv_text.split())

    cv_skill_list = sorted(cv_skills)
    matched_list = sorted(matched_required)
    missing_required_list = sorted(missing_required)
    missing_preferred_list = sorted(missing_preferred)

    cv_skill_text = (
        ", ".join(cv_skill_list)
        if cv_skill_list
        else "No taxonomy skills confidently detected"
    )

    if cv_skill_list:
        summary = (
            "Professional summary draft based on the uploaded CV. "
            f"Detected skills include: {cv_skill_text}. "
            "Personalize this summary with verified education, "
            "projects, and experience from the CV. "
            "Skill detection alone does not confirm proficiency."
        )
    else:
        summary = (
            "Professional summary draft: Add verified education, "
            "projects, skills, and career goals from the uploaded CV."
        )

    bullets = []

    # Use a short excerpt from the actual CV as evidence.
    if cv_text_clean:
        excerpt_length = 500
        cv_excerpt = cv_text_clean[:excerpt_length]

        if len(cv_text_clean) > excerpt_length:
            cv_excerpt += "..."

        bullets.append(
            "CV evidence excerpt for review: "
            f"{cv_excerpt} "
            "Use the original CV to verify project names, "
            "responsibilities, and results before creating final bullets."
        )

    if matched_list:
        bullets.append(
            "Relevant detected skills: "
            + ", ".join(matched_list)
            + ". Connect these skills to specific CV-documented "
            "projects or tasks before adding them to a resume."
        )
    else:
        bullets.append(
            "Relevant detected skills: No required skills were "
            "confidently matched. Review the CV and job description."
        )

    if missing_required_list:
        bullets.append(
            "Required skills to review: "
            + ", ".join(missing_required_list)
            + ". These were not detected in the CV; this does not "
            "prove the candidate lacks them."
        )

    if missing_preferred_list:
        bullets.append(
            "Preferred skills to consider: "
            + ", ".join(missing_preferred_list)
            + ". Include these only if they accurately represent "
            "the candidate's abilities."
        )

    bullets.append(
        "Evidence reminder: Do not invent projects, responsibilities, "
        "experience, or achievements."
    )

    bullets.append(
        "Matching note: TF-IDF text similarity is "
        f"{similarity_score}%. It measures text overlap, not hiring "
        "probability or verified job suitability."
    )

    return summary, bullets



# ==================================================
# 8. WORD DOCUMENT EXPORT
# ==================================================

def create_docx_file(summary_text, bullet_points):
    document = Document()

    document.add_heading(
        "CareerForge AI - Resume Content", level=1
    )
    document.add_heading("Professional Summary Draft", level=2)
    document.add_paragraph(summary_text)

    document.add_heading("Skills and Review Notes", level=2)

    for bullet in bullet_points:
        document.add_paragraph(bullet, style="List Bullet")

    document.add_paragraph(
        "Review all generated content before using it in a CV."
    )

    file_buffer = io.BytesIO()
    document.save(file_buffer)
    file_buffer.seek(0)

    return file_buffer.getvalue()


# ==================================================
# 9. SIDEBAR
# ==================================================

with st.sidebar:
    st.header("About CareerForge AI")

    st.write(
        "CareerForge AI compares CVs with job descriptions "
        "using a reusable skill taxonomy and TF-IDF text similarity."
    )

    st.caption(
        "This is an academic prototype. Detected skills and scores "
        "should be reviewed and are not proof of proficiency."
    )


# ==================================================
# 10. UPLOAD AND PROCESS CV
# ==================================================

st.header("1. Upload Candidate Resume")

uploaded_file = st.file_uploader(
    "Choose your CV or resume",
    type=["pdf", "docx", "txt"]
)

cv_text = ""
cv_skills_by_category = {}
cv_skills = set()

if uploaded_file is not None:
    try:
        cv_text = extract_text(uploaded_file)

        st.success("CV text extracted successfully.")

        with st.expander("Preview extracted CV text"):
            st.text(cv_text[:10000])

        cv_skills_by_category, cv_skills = extract_all_skills(cv_text)

        st.subheader("Skills Detected in Your CV")

        if cv_skills_by_category:
            for category, skills in cv_skills_by_category.items():
                st.markdown(f"**{category}**")
                st.write(", ".join(skills))
        else:
            st.info("No skills from the current taxonomy were detected.")

        st.caption(
            "Verify detected skills against the original CV. "
            "A mention does not establish proficiency."
        )

    except Exception as e:
        cv_text = ""
        cv_skills_by_category = {}
        cv_skills = set()

        st.error(f"Could not read the uploaded file: {e}")


# ==================================================
# 11. JOB DESCRIPTION
# ==================================================

st.header("2. Target Job Description")

job_description = st.text_area(
    "Paste the job description here",
    height=250,
    placeholder="Paste the complete job description here..."
)


# ==================================================
# 12. ANALYSIS
# ==================================================

if st.button("Analyze CV and Job Description", type="primary"):

    if uploaded_file is None:
        st.error("Please upload your CV first.")

    elif not cv_text.strip():
        st.error("No readable CV text was found.")

    elif not job_description.strip():
        st.error("Please paste a job description.")

    else:
        required_skills, preferred_skills = (
            extract_requirement_sections(job_description)
        )

        matched_required = cv_skills.intersection(required_skills)
        missing_required = required_skills.difference(cv_skills)

        matched_preferred = cv_skills.intersection(preferred_skills)
        missing_preferred = preferred_skills.difference(cv_skills)

        matched_skills = matched_required.union(matched_preferred)
        missing_skills = missing_required.union(missing_preferred)

        if required_skills:
            taxonomy_score = round(
                len(matched_required) / len(required_skills) * 100
            )
        else:
            taxonomy_score = 0

        similarity_score = calculate_text_similarity(
            cv_text,
            job_description
        )

        
        summary_text, bullet_points = generate_resume_content(
           cv_text,
           cv_skills,
           matched_required,
           missing_required,
           matched_preferred,
           missing_preferred,
            similarity_score
          )


        st.session_state["analysis_result"] = {
            "cv_skills": sorted(cv_skills),
            "required_skills": sorted(required_skills),
            "preferred_skills": sorted(preferred_skills),
            "matched_required": sorted(matched_required),
            "missing_required": sorted(missing_required),
            "matched_preferred": sorted(matched_preferred),
            "missing_preferred": sorted(missing_preferred),
            "matched_skills": sorted(matched_skills),
            "missing_skills": sorted(missing_skills),
            "taxonomy_score": taxonomy_score,
            "similarity_score": similarity_score,
            "summary_text": summary_text,
            "bullet_points": bullet_points,
        }


# ==================================================
# 13. DISPLAY RESULTS
# ==================================================

if "analysis_result" in st.session_state:
    result = st.session_state["analysis_result"]

    st.divider()
    st.header("3. Analysis Results")

    col1, col2 = st.columns(2)

    with col1:
        st.metric(
            "TF-IDF Text Similarity",
            f"{result['similarity_score']}%"
        )

    with col2:
        st.metric(
            "Required-Skill Coverage",
            f"{result['taxonomy_score']}%"
        )

    st.progress(result["taxonomy_score"] / 100)

    st.caption(
        "Required-skill coverage is based only on skills detected "
        "by the current taxonomy. It does not measure every job "
        "requirement or predict the chance of being hired."
    )

    # Required Skills
    st.subheader("Required Skills")

    st.markdown("**Matched required skills**")
    if result["matched_required"]:
        st.write(", ".join(result["matched_required"]))
    else:
        st.info("No required skills matched.")

    st.markdown("**Required skills to review**")
    if result["missing_required"]:
        st.write(", ".join(result["missing_required"]))
    else:
        st.success("All detected required skills matched.")

    # Preferred Skills
    st.subheader("Preferred Skills")

    st.markdown("**Matched preferred skills**")
    if result["matched_preferred"]:
        st.write(", ".join(result["matched_preferred"]))
    else:
        st.write("No preferred skills matched.")

    st.markdown("**Preferred skills to consider learning**")
    if result["missing_preferred"]:
        st.write(", ".join(result["missing_preferred"]))
    elif result["preferred_skills"]:
        st.write("All detected preferred skills matched.")
    else:
        st.write(
            "No preferred skills were identified in the job description."
        )

    # All detected job-description skills
    with st.expander("View all detected job-description skills"):
        st.markdown("**Required skills**")
        st.write(
            ", ".join(result["required_skills"])
            if result["required_skills"]
            else "None detected"
        )

        st.markdown("**Preferred skills**")
        st.write(
            ", ".join(result["preferred_skills"])
            if result["preferred_skills"]
            else "None detected"
        )

    # Generated Resume Content
    st.divider()
    st.header("4. Generated Resume Content")

    st.subheader("Professional Summary Draft")
    st.write(result["summary_text"])

    st.subheader("Resume Bullet Points")
    for bullet in result["bullet_points"]:
        st.markdown(f"- {bullet}")

    st.warning(
        "Check all generated content before using it in a CV. "
        "Only include skills and experience you genuinely possess."
    )

    # Prepare downloadable text file
    text_content = (
        "CAREERFORGE AI - RESUME CONTENT\n\n"
        "PROFESSIONAL SUMMARY\n"
        f"{result['summary_text']}\n\n"
        "SKILLS AND REVIEW NOTES\n"
        + "\n".join(
            f"- {bullet}"
            for bullet in result["bullet_points"]
        )
    )

    # Prepare downloadable Word document
    docx_content = create_docx_file(
        result["summary_text"],
        result["bullet_points"]
    )

    download_col1, download_col2 = st.columns(2)

    with download_col1:
        st.download_button(
            label="Download Word Document",
            data=docx_content,
            file_name="CareerForge_Resume_Content.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

    with download_col2:
        st.download_button(
            label="Download Text File",
            data=text_content,
            file_name="CareerForge_Resume_Content.txt",
            mime="text/plain",
        )


# ==================================================
# 14. FOOTER
# ==================================================

st.divider()
st.caption("CareerForge AI | Academic Project Prototype")
